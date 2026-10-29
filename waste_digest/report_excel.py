"""Filterable Excel version of the digest."""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F6F4A")
NEW_FILL = PatternFill("solid", fgColor="FDEEE5")


def _sheet(ws, headers, widths):
    ws.append(headers)
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"


def _dt(iso):
    if not iso:
        return None
    return datetime.fromisoformat(iso).replace(tzinfo=None)  # Excel has no time zones; values are UTC


def write(digest: dict, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Digest"
    headers = ["New?", "Chatter?", "Section", "Companies", "Topics", "Headline", "Source", "Published (UTC)",
               "Summary", "# Outlets", "Other outlets", "Link"]
    _sheet(ws, headers, [7, 9, 26, 14, 30, 70, 20, 17, 70, 9, 40, 40])
    wrap = Alignment(wrap_text=True, vertical="top")
    for s in digest["sections"]:
        for i in s["items"]:
            ws.append(["NEW" if i["is_new"] else "", "Y" if i.get("chatter") else "", s["title"], ", ".join(i["companies"]),
                       "; ".join(i["topic_names"]), i["headline"], i["source"], _dt(i["published"]),
                       i["summary"], 1 + len(i["others"]), ", ".join(o["source"] for o in i["others"]), i["url"]])
            r = ws.max_row
            ws.cell(r, 6).hyperlink = i["url"]
            ws.cell(r, 6).font = Font(color="14532D", underline="single")
            ws.cell(r, 8).number_format = "yyyy-mm-dd hh:mm"
            for col in (6, 9, 11):
                ws.cell(r, col).alignment = wrap
            if i["is_new"]:
                ws.cell(r, 1).fill = NEW_FILL
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("All articles")
    _sheet(ws2, ["Story headline", "Outlet", "Article headline", "Published (local)", "Link"], [60, 22, 60, 18, 50])
    for s in digest["sections"]:
        for i in s["items"]:
            ws2.append([i["headline"], i["source"], i["headline"], i["published_local"], i["url"]])
            for o in i["others"]:
                ws2.append([i["headline"], o["source"], o["headline"], o["published_local"], o["url"]])
    ws2.auto_filter.ref = ws2.dimensions

    ws3 = wb.create_sheet("Run info")
    _sheet(ws3, ["Item", "Value"], [28, 90])
    ws3.append(["Generated", digest["generated_local"]])
    ws3.append(["Lookback (hours)", digest["hours"]])
    for k, v in digest["stats"].items():
        ws3.append([k.replace("_", " ").capitalize(), v])
    for src in digest["sources"]:
        ws3.append([src["name"], f'{src["items"]} items; {src["failures"]} of {src["requests"]} requests failed'
                    + ("; " + " | ".join(src["errors"][:3]) if src["errors"] else "")])
    wb.save(path)
